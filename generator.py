from __future__ import annotations

import html
import os
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from config import FONT_NAME, FOOTER_TEXT, LOGO_PATH


def sanitize_text(text: str) -> str:
    """Normalize user text for preview, DOCX, and PDF export."""
    if text is None:
        return ""

    cleaned = str(text)
    cleaned = cleaned.replace("\u00A0", " ")
    cleaned = cleaned.replace("“", '"').replace("”", '"')
    cleaned = cleaned.replace("‘", "'").replace("’", "'")
    cleaned = cleaned.replace("–", "-").replace("—", "-")
    cleaned = cleaned.replace("…", "...")
    cleaned = cleaned.replace("•", "- ").replace("▪", "- ")
    cleaned = cleaned.replace("•", "- ").replace("▪", "- ")
    cleaned = cleaned.replace("**", "")
    cleaned = re.sub(r"(?m)^\s*#+\s*", "", cleaned)
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"\*\*+", "", cleaned)
    cleaned = re.sub(r"(?m)^\s*[-*]\s*", "- ", cleaned)
    cleaned = cleaned.encode("latin-1", "ignore").decode("latin-1")
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    """Render the legal document as a DOCX file bytes object with branding."""
    sanitized = sanitize_text(text)
    document = Document()

    default_style = document.styles["Normal"]
    default_style.font.name = FONT_NAME
    default_style.font.size = Pt(12)

    if os.path.exists(LOGO_PATH):
        try:
            document.add_paragraph()
            document.add_picture(LOGO_PATH, width=Inches(2.0))
        except Exception:
            pass

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run((doc_type or "Legal Document").strip() or "Legal Document")
    title_run.bold = True
    title_run.font.name = FONT_NAME
    title_run.font.size = Pt(18)

    if terms and terms.strip():
        concept_paragraph = document.add_paragraph()
        concept_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        concept_run = concept_paragraph.add_run("Terms and Conditions")
        concept_run.bold = True
        concept_run.font.name = FONT_NAME
        concept_run.font.size = Pt(12)

        rows = [part.strip() for part in str(terms).split(";") if part and part.strip()]
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"

        for index, item in enumerate(rows, start=1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(index)
            row_cells[1].text = item

    for raw_line in sanitized.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
            continue

        is_heading = bool(
            re.match(r"^\d+\.\s+[A-Za-z0-9].*:$", line)
            or (line.endswith(":") and (line.isupper() or len(line) < 80))
        )

        if is_heading:
            heading_paragraph = document.add_paragraph()
            heading_run = heading_paragraph.add_run(line)
            heading_run.bold = True
            heading_run.font.name = FONT_NAME
            heading_run.font.size = Pt(12)
            continue

        if line.startswith(("- ", "* ", "•")):
            bullet_paragraph = document.add_paragraph(style="List Bullet")
            bullet_paragraph.add_run(line.lstrip("-*• "))
            continue

        body_paragraph = document.add_paragraph(line)
        body_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        body_paragraph.style = default_style

    for section in document.sections:
        footer_paragraph = section.footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_paragraph.text = FOOTER_TEXT
        footer_run = footer_paragraph.runs[0]
        footer_run.italic = True
        footer_run.font.name = FONT_NAME
        footer_run.font.size = Pt(8)
        footer_paragraph.add_run("  ")
        page_run = footer_paragraph.add_run("Page 1")
        page_run.font.name = FONT_NAME
        page_run.font.size = Pt(8)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class LegalPDF(FPDF):
    """Custom PDF subclass for LegalEase branding and page layout."""

    def __init__(self, title: str):
        super().__init__()
        self.title = title

    def header(self) -> None:
        self.set_font("Helvetica", "B", 12)
        self.set_y(8)
        if os.path.exists(LOGO_PATH):
            try:
                self.image(LOGO_PATH, x=(self.w - 30) / 2, y=8, w=25)
                self.ln(16)
            except Exception:
                pass
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 12, self.title, 0, 1, "C")
        self.ln(2)

    def footer(self) -> None:
        self.set_y(-20)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 8, FOOTER_TEXT, 0, 0, "C")
        self.set_x(self.w - 52)
        self.cell(0, 8, f"Page {self.page_no()} of {{nb}}", 0, 0, "R")


def format_pdf(text: str, doc_type: str) -> bytes:
    """Render a branded PDF containing the generated legal document."""
    sanitized = sanitize_text(text)
    pdf = LegalPDF(doc_type or "Legal Document")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("Helvetica", "", 11)

    for raw_line in sanitized.splitlines():
        line = raw_line.strip()
        if not line:
            pdf.ln(4)
            continue

        if re.match(r"^\d+\.\s+[A-Za-z0-9].*:$", line) or (line.endswith(":") and (line.isupper() or len(line) < 80)):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 8, line)
            pdf.ln(2)
            continue

        if line.startswith(("- ", "* ", "•")):
            pdf.set_font("Helvetica", "", 11)
            pdf.cell(6, 8, "-", 0, 0)
            pdf.multi_cell(0, 8, line.lstrip("-*• "))
            pdf.ln(2)
            continue

        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 8, line)
        pdf.ln(2)

    result = pdf.output(dest="S")
    if isinstance(result, str):
        return result.encode("latin-1", "ignore")
    if result is None:
        return b""
    return bytes(result)


def format_html_preview(text: str) -> str:
    """Convert the legal text to a dark themed HTML preview."""
    sanitized = sanitize_text(text)
    lines = sanitized.splitlines()
    blocks: list[str] = []
    bullet_items: list[str] = []

    def flush_bullets() -> None:
        if not bullet_items:
            return
        blocks.append("<ul style='margin: 0 0 12px 18px; padding: 0; color: #F3F4F6;'>")
        for item in bullet_items:
            blocks.append(f"<li style='margin-bottom: 6px;'>{item}</li>")
        blocks.append("</ul>")
        bullet_items.clear()

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            flush_bullets()
            continue

        is_heading = bool(
            re.match(r"^\d+\.\s+[A-Za-z0-9].*:$", line)
            or (line.endswith(":") and (line.isupper() or len(line) < 80))
        )

        if is_heading:
            flush_bullets()
            blocks.append(
                f"<h4 style='color: #F9FAFB; margin: 18px 0 8px; font-size: 1.02rem; font-weight: 700;'>{html.escape(line)}</h4>"
            )
            continue

        if line.startswith(("- ", "* ", "•")):
            bullet_items.append(html.escape(line.lstrip("-*• ")))
            continue

        flush_bullets()
        blocks.append(
            f"<p style='margin: 0 0 10px; line-height: 1.6; color: #E5E7EB;'>{html.escape(line)}</p>"
        )

    flush_bullets()
    content = "".join(blocks)
    return (
        "<div style='background: #111827; color: #F9FAFB; border-radius: 12px; padding: 16px; "
        "overflow-y: auto; max-height: 400px; font-family: Arial, sans-serif; line-height: 1.5;'>"
        f"{content}"
        "</div>"
    )
